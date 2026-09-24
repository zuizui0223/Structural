#!/usr/bin/env python3
"""Freeze Roeble et al. 2024 geological-origin moderator for the community panel."""
from __future__ import annotations

from collections import defaultdict
import argparse, hashlib, io, json
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

URL="https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-024-51556-7/MediaObjects/41467_2024_51556_MOESM6_ESM.xlsx"
EXPECTED_SHA256="fdca6086440c721e2766c3b73b96969c401b1a267a0cf9590c912d369d5e8597"
DOI="10.1038/s41467-024-51556-7"
MIN_TYPE_COVERAGE=0.80

def sha(x):
    return hashlib.sha256(
        json.dumps(x,sort_keys=True,separators=(",",":"),ensure_ascii=False).encode()
    ).hexdigest()

def fetch_bytes(url):
    req=Request(url,headers={"User-Agent":"Structural-GIFT-community-geology/0.1"})
    with urlopen(req,timeout=180) as r:return r.read()

def norm_type(x):
    if x is None:return None
    t=str(x).strip().lower()
    return t if t in {"continental","oceanic"} else None

def main():
    ap=argparse.ArgumentParser();ap.add_argument("--panel",type=Path,required=True);a=ap.parse_args()
    panel=json.loads(a.panel.read_text())
    if panel.get("schema")!="structural.gift_community_topology_panel.v0_2":raise RuntimeError("unexpected panel")
    if panel.get("response_values_accessed") is not False:raise RuntimeError("panel response opened")

    data=fetch_bytes(URL); digest=hashlib.sha256(data).hexdigest()
    if digest!=EXPECTED_SHA256:raise RuntimeError(f"Roeble supplement SHA drift: {digest}")
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True);ws=wb[wb.sheetnames[0]]
    rows=ws.iter_rows(values_only=True);header=[str(x) if x is not None else "" for x in next(rows)]
    required=("entity_ID","island_type")
    missing=[x for x in required if x not in header]
    if missing:raise RuntimeError("missing columns: "+",".join(missing))
    ix={x:header.index(x) for x in required}
    by_entity=defaultdict(list)
    for row in rows:
        try:eid=str(int(float(row[ix["entity_ID"]])))
        except (TypeError,ValueError):continue
        t=norm_type(row[ix["island_type"]])
        if t:by_entity[eid].append(t)

    canonical={}
    conflicts=[]
    for eid,vals in by_entity.items():
        types=sorted(set(vals))
        if len(types)==1:canonical[eid]=types[0]
        elif len(types)>1:conflicts.append({"entity_ID":eid,"types":types})

    groups=[]
    for g in panel["groups"]:
        ids=[str(i["entity_ID"]) for i in g["islands"]]
        typed=[canonical[eid] for eid in ids if eid in canonical]
        coverage=len(typed)/len(ids)
        frac=(sum(t=="oceanic" for t in typed)/len(typed)) if typed and coverage>=MIN_TYPE_COVERAGE else None
        groups.append({
            "archipelago_id":g["archipelago_id"],
            "n_islands":len(ids),
            "typed_islands":len(typed),
            "type_coverage":coverage,
            "oceanic_fraction":frac,
            "geology_class":(
                None if frac is None else
                "continental" if frac==0 else
                "oceanic" if frac==1 else
                "mixed"
            ),
        })

    eligible=[g for g in groups if g["oceanic_fraction"] is not None]
    vals=[g["oceanic_fraction"] for g in eligible]
    payload={
        "schema":"structural.gift_community_roeble_geology.v0_1",
        "status":"FROZEN_BEFORE_COMMUNITY_RESPONSE",
        "panel_fingerprint":panel["panel_fingerprint"],
        "response_values_accessed":False,
        "gift_species_composition_accessed":False,
        "source":{
            "doi":DOI,"supplement":"Supplementary Data 3","url":URL,
            "xlsx_sha256":digest,"xlsx_bytes":len(data),
        },
        "matching_rule":{
            "key":"exact GIFT entity_ID",
            "minimum_within_archipelago_typed_fraction":MIN_TYPE_COVERAGE,
            "no_name_fuzzy_matching":True,
            "no_response_based_resolution":True,
        },
        "duplicate_entity_type_conflicts":conflicts,
        "groups":groups,
        "eligible_archipelagos":len(eligible),
        "oceanic_fraction_min":min(vals) if vals else None,
        "oceanic_fraction_max":max(vals) if vals else None,
        "pure_continental":sum(v==0 for v in vals),
        "pure_oceanic":sum(v==1 for v in vals),
        "mixed":sum(0<v<1 for v in vals),
        "continuous_H2_support":len(eligible)>=6 and min(vals)<max(vals) if vals else False,
        "claim_boundary":"geological origin/past land connection moderator only; not direct evidence of colonization, persistence or extinction mechanism",
    }
    payload["crosswalk_fingerprint"]=sha({
        "source_sha256":digest,"matching_rule":payload["matching_rule"],"groups":groups,
    })
    print(json.dumps(payload,indent=2,sort_keys=True));return 0

if __name__=="__main__":raise SystemExit(main())
