#!/usr/bin/env python3
"""Freeze Roeble cleaned-archipelago geological moderator for community H2."""
from __future__ import annotations

from collections import Counter, defaultdict
import argparse, hashlib, io, json
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

URL="https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-024-51556-7/MediaObjects/41467_2024_51556_MOESM6_ESM.xlsx"
EXPECTED_SHA256="fdca6086440c721e2766c3b73b96969c401b1a267a0cf9590c912d369d5e8597"
DOI="10.1038/s41467-024-51556-7"
MIN_ENTITY_MATCHES=5
MIN_ENTITY_COVERAGE=0.50
MIN_DOMINANT_ARCHIPELAGO_SHARE=0.80

def sha(x):
    return hashlib.sha256(
        json.dumps(x,sort_keys=True,separators=(",",":"),ensure_ascii=False).encode()
    ).hexdigest()

def fetch_bytes(url):
    req=Request(url,headers={"User-Agent":"Structural-GIFT-community-geology-archipelago/0.1"})
    with urlopen(req,timeout=180) as r:return r.read()

def norm_type(x):
    if x is None:return None
    t=str(x).strip().lower()
    return t if t in {"continental","oceanic"} else None

def main():
    ap=argparse.ArgumentParser();ap.add_argument("--panel",type=Path,required=True);a=ap.parse_args()
    panel=json.loads(a.panel.read_text())
    if panel.get("schema")!="structural.gift_community_topology_panel.v0_2":raise RuntimeError("unexpected panel")
    if panel.get("response_values_accessed") is not False:raise RuntimeError("panel response opened")

    data=fetch_bytes(URL); digest=hashlib.sha256(data).hexdigest()
    if digest!=EXPECTED_SHA256:raise RuntimeError(f"Roeble supplement SHA drift: {digest}")
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True);ws=wb[wb.sheetnames[0]]
    rows=ws.iter_rows(values_only=True);header=[str(x) if x is not None else "" for x in next(rows)]
    required=("entity_ID","archipelago","island_type")
    missing=[x for x in required if x not in header]
    if missing:raise RuntimeError("missing columns: "+",".join(missing))
    ix={x:header.index(x) for x in required}

    raw=[]
    for row in rows:
        try:eid=str(int(float(row[ix["entity_ID"]])))
        except (TypeError,ValueError):continue
        arch=None if row[ix["archipelago"]] is None else str(row[ix["archipelago"]]).strip()
        typ=norm_type(row[ix["island_type"]])
        raw.append({"entity_ID":eid,"archipelago":arch,"island_type":typ})

    by_entity=defaultdict(list); by_arch=defaultdict(list)
    for row in raw:
        by_entity[row["entity_ID"]].append(row)
        if row["archipelago"]:by_arch[row["archipelago"]].append(row)

    canonical={}; conflicts=[]
    for eid,vals in by_entity.items():
        pairs={(v["archipelago"],v["island_type"]) for v in vals}
        if len(pairs)==1:canonical[eid]=vals[0]
        elif len(pairs)>1:conflicts.append({"entity_ID":eid,"pairs":sorted([list(p) for p in pairs],key=str)})

    arch_types={}
    for name,vals in by_arch.items():
        typed=[v["island_type"] for v in vals if v["island_type"]]
        c=Counter(typed)
        denom=sum(c.values())
        arch_types[name]={
            "typed_rows":denom,
            "continental":c["continental"],
            "oceanic":c["oceanic"],
            "oceanic_fraction":c["oceanic"]/denom if denom else None,
            "class":(
                "continental" if denom and c["oceanic"]==0 else
                "oceanic" if denom and c["continental"]==0 else
                "mixed" if denom else None
            ),
        }

    groups=[]
    for g in panel["groups"]:
        ids=[str(i["entity_ID"]) for i in g["islands"]]
        matched=[canonical[eid] for eid in ids if eid in canonical]
        names=Counter(r["archipelago"] for r in matched if r["archipelago"])
        dominant,dominant_n=(None,0)
        if names:
            dominant,dominant_n=sorted(names.items(),key=lambda kv:(-kv[1],kv[0]))[0]
        match_fraction=len(matched)/len(ids)
        dominant_share=dominant_n/len(matched) if matched else 0.0
        qualified=(
            dominant is not None
            and len(matched)>=MIN_ENTITY_MATCHES
            and match_fraction>=MIN_ENTITY_COVERAGE
            and dominant_share>=MIN_DOMINANT_ARCHIPELAGO_SHARE
            and arch_types.get(dominant,{}).get("oceanic_fraction") is not None
        )
        external=arch_types.get(dominant) if qualified else None
        groups.append({
            "archipelago_id":g["archipelago_id"],
            "n_islands":len(ids),
            "entity_matches":len(matched),
            "entity_match_fraction":match_fraction,
            "dominant_roeble_archipelago":dominant,
            "dominant_roeble_archipelago_matches":dominant_n,
            "dominant_share_among_matches":dominant_share,
            "mapping_qualified":qualified,
            "geology_class":external["class"] if external else None,
            "oceanic_fraction":external["oceanic_fraction"] if external else None,
        })

    eligible=[g for g in groups if g["oceanic_fraction"] is not None]
    vals=[g["oceanic_fraction"] for g in eligible]
    payload={
        "schema":"structural.gift_community_roeble_geology.v0_2",
        "status":"FROZEN_BEFORE_COMMUNITY_RESPONSE",
        "panel_fingerprint":panel["panel_fingerprint"],
        "response_values_accessed":False,
        "gift_species_composition_accessed":False,
        "source":{
            "doi":DOI,"supplement":"Supplementary Data 3","url":URL,
            "xlsx_sha256":digest,"xlsx_bytes":len(data),
        },
        "matching_rule":{
            "entity_key":"exact GIFT entity_ID",
            "minimum_entity_matches":MIN_ENTITY_MATCHES,
            "minimum_entity_match_fraction":MIN_ENTITY_COVERAGE,
            "minimum_dominant_cleaned_archipelago_share":MIN_DOMINANT_ARCHIPELAGO_SHARE,
            "moderator":"oceanic fraction across all typed islands in the matched Roeble cleaned archipelago",
            "no_fuzzy_name_matching":True,
            "no_response_based_resolution":True,
        },
        "duplicate_entity_conflicts":conflicts,
        "groups":groups,
        "eligible_archipelagos":len(eligible),
        "oceanic_fraction_min":min(vals) if vals else None,
        "oceanic_fraction_max":max(vals) if vals else None,
        "pure_continental":sum(v==0 for v in vals),
        "pure_oceanic":sum(v==1 for v in vals),
        "mixed":sum(0<v<1 for v in vals),
        "continuous_H2_support":len(eligible)>=6 and min(vals)<max(vals) if vals else False,
        "claim_boundary":"published cleaned-archipelago geological-origin moderator; not a demographic process estimate",
    }
    payload["crosswalk_fingerprint"]=sha({
        "source_sha256":digest,"matching_rule":payload["matching_rule"],"groups":groups,
    })
    print(json.dumps(payload,indent=2,sort_keys=True));return 0

if __name__=="__main__":raise SystemExit(main())
