#!/usr/bin/env python3
"""One-shot execution of the frozen murundu relative-topology protocol.

The Dryad workbook may be opened only after an exact authorization receipt is
committed. Geometry columns are parsed first. Richness.Trees is parsed only if
the frozen geometry gate passes. No herb or termite response is read.
"""
from __future__ import annotations

from collections import defaultdict
import argparse
import hashlib
import json
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

EXPECTED_PROTOCOL="83b6a2539acc7983bb75648eb6c53a913917d5ea50dffdaefceceb3ac8fb6512"
EXPECTED_PLOTS=tuple(f"CM{i}" for i in range(1,12))
GEOM_COLUMNS=("ID.murundu","Plots","X","Y","Area_Mur_m2","Height_Mur")
PRIMARY_RESPONSE="Richness.Trees"

def sha(value)->str:
    return hashlib.sha256(
        json.dumps(value,sort_keys=True,separators=(",",":"),ensure_ascii=False).encode()
    ).hexdigest()

def load_json(path:Path)->dict:
    x=json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(x,dict):
        raise RuntimeError(f"{path} must contain a JSON object")
    return x

def file_sha(path:Path)->str:
    h=hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda:f.read(1024*1024),b""):
            h.update(chunk)
    return h.hexdigest()

def number(x):
    if x is None or x=="":
        return None
    try:
        v=float(x)
    except (TypeError,ValueError):
        return None
    return v if math.isfinite(v) else None

def haversine_m(lon1,lat1,lon2,lat2):
    r=6371008.8
    p1,p2=math.radians(lat1),math.radians(lat2)
    dp=p2-p1
    dl=math.radians(lon2-lon1)
    a=math.sin(dp/2)**2+math.cos(p1)*math.cos(p2)*math.sin(dl/2)**2
    return 2*r*math.asin(min(1.0,math.sqrt(a)))

def col_values(ws,col_idx,nrows):
    vals=[]
    for row in ws.iter_rows(
        min_row=2,max_row=nrows+1,min_col=col_idx,max_col=col_idx,values_only=True
    ):
        vals.append(row[0])
    if len(vals)!=nrows:
        raise RuntimeError(f"column {col_idx}: expected {nrows} rows, got {len(vals)}")
    return vals

def minimax_matrix(nodes):
    n=len(nodes)
    d=np.zeros((n,n),dtype=float)
    for i in range(n):
        for j in range(i+1,n):
            x=haversine_m(nodes[i]["lon"],nodes[i]["lat"],nodes[j]["lon"],nodes[j]["lat"])
            d[i,j]=d[j,i]=x
    b=d.copy()
    # Floyd-Warshall on the minimax semiring.
    for k in range(n):
        b=np.minimum(b,np.maximum(b[:,k,None],b[None,k,:]))
    return d,b

def z(values,label):
    a=np.asarray(values,dtype=float)
    mu=float(a.mean())
    sd=float(a.std(ddof=0))
    if sd<=1e-12:
        raise RuntimeError(f"constant predictor {label}")
    return (a-mu)/sd,{"mean":mu,"sd":sd}

def design_for_plot(focals):
    za,s_area=z([math.log(r["area"]) for r in focals],"log_area")
    zh,s_height=z([math.log1p(r["height"]) for r in focals],"log1p_height")
    zd,s_direct=z([math.log1p(r["direct_m"]) for r in focals],"log1p_direct")
    zg,s_gain=z([r["gain"] for r in focals],"topology_gain")
    inter=zd*zg
    X=np.column_stack([np.ones(len(focals)),za,zh,zd,zg,inter])
    for i,r in enumerate(focals):
        r["design"]=[float(x) for x in X[i,:]]
    rank=int(np.linalg.matrix_rank(X,tol=1e-10))
    sv=np.linalg.svd(X,compute_uv=False)
    cond=float(sv[0]/sv[-1]) if sv[-1]>1e-15 else float("inf")
    return X,{
        "rank":rank,
        "n_columns":int(X.shape[1]),
        "condition_number":cond,
        "full_rank":rank==X.shape[1],
        "scaling":{
            "log_area":s_area,
            "log1p_height":s_height,
            "log1p_direct_source_isolation":s_direct,
            "topology_gain":s_gain,
        },
    }

def geometry_stage(records,protocol):
    by=defaultdict(list)
    missing=0
    for r in records:
        if (
            not r["plot"] or r["lon"] is None or r["lat"] is None
            or r["area"] is None or r["area"]<=0
        ):
            missing+=1
            continue
        by[r["plot"]].append(r)

    declared=set(protocol["selection_provenance"]["declared_plots"])
    observed=set(by)
    if not observed.issubset(declared):
        raise RuntimeError(f"unexpected plot IDs: {sorted(observed-declared)}")

    gate=protocol["geometry_gate"]["per_plot"]
    audits=[]
    qualified={}
    total_focals=0
    for plot in sorted(declared,key=lambda x:int(x[2:])):
        nodes=by.get(plot,[])
        if not nodes:
            audits.append({"plot":plot,"pass":False,"reason":"no geometry nodes"})
            continue
        d,b=minimax_matrix(nodes)
        focals=[]
        for i,r in enumerate(nodes):
            sources=[j for j,s in enumerate(nodes) if s["area"]>r["area"]]
            if not sources:
                continue
            if r["height"] is None or r["height"]<0:
                continue
            direct=min(float(d[i,j]) for j in sources)
            bottleneck=min(float(b[i,j]) for j in sources)
            gain=max(0.0,math.log1p(direct)-math.log1p(bottleneck))
            focals.append({
                "row_index":r["row_index"],
                "id":r["id"],
                "area":r["area"],
                "height":r["height"],
                "direct_m":direct,
                "bottleneck_m":bottleneck,
                "gain":gain,
            })
        positive=sum(r["gain"]>1e-12 for r in focals)
        audit={
            "plot":plot,
            "graph_nodes":len(nodes),
            "focal_murundus":len(focals),
            "positive_topology_gain":positive,
        }
        try:
            X,da=design_for_plot(focals) if focals else (None,None)
        except RuntimeError as exc:
            da=None
            audit["design_error"]=str(exc)
        if da:
            audit["design_audit"]=da
        passed=(
            len(focals)>=int(gate["minimum_focal_murundus"])
            and positive>=int(gate["minimum_positive_topology_gain"])
            and da is not None
            and da["full_rank"]
            and da["condition_number"]<=float(gate["maximum_condition_number"])
        )
        audit["pass"]=bool(passed)
        audits.append(audit)
        if passed:
            qualified[plot]=focals
            total_focals+=len(focals)

    study_gate=protocol["geometry_gate"]["study_pass"]
    study_pass=(
        len(qualified)>=int(study_gate["minimum_geometry_qualified_plots"])
        and total_focals>=int(study_gate["minimum_total_geometry_qualified_focal_murundus"])
    )
    return qualified,{
        "missing_geometry_rows":missing,
        "plot_audits":audits,
        "qualified_plots":sorted(qualified,key=lambda x:int(x[2:])),
        "qualified_plot_count":len(qualified),
        "qualified_focal_murundus":total_focals,
        "study_pass":bool(study_pass),
    }

def response_stage(qualified,response_values,protocol):
    gate=protocol["response_gate"]["per_plot"]
    target_index=protocol["per_plot_model"]["columns"].index(
        protocol["per_plot_model"]["target_column"]
    )
    direct_index=protocol["per_plot_model"]["columns"].index(
        "z_log1p_direct_source_isolation"
    )
    gain_index=protocol["per_plot_model"]["columns"].index("z_topology_gain")
    audits=[]
    coefs={}
    for plot in sorted(qualified,key=lambda x:int(x[2:])):
        focals=qualified[plot]
        keep=[]
        invalid_negative=0
        for r in focals:
            y=number(response_values[r["row_index"]])
            if y is None:
                continue
            if y<0:
                invalid_negative+=1
                continue
            keep.append((r,y))
        fraction=len(keep)/len(focals) if focals else 0.0
        audit={
            "plot":plot,
            "geometry_focals":len(focals),
            "nonmissing_response_rows":len(keep),
            "nonmissing_fraction":fraction,
            "invalid_negative_response_rows":invalid_negative,
        }
        if keep:
            X=np.asarray([r["design"] for r,_ in keep],dtype=float)
            ylog=np.log1p(np.asarray([y for _,y in keep],dtype=float))
            ysd=float(ylog.std(ddof=0))
            audit["response_log1p_sd"]=ysd
            rank=int(np.linalg.matrix_rank(X,tol=1e-10))
            sv=np.linalg.svd(X,compute_uv=False)
            cond=float(sv[0]/sv[-1]) if sv[-1]>1e-15 else float("inf")
            audit["design_after_missing"]={
                "rank":rank,"n_columns":int(X.shape[1]),
                "full_rank":rank==X.shape[1],"condition_number":cond,
                "target_retained":rank==X.shape[1],
            }
            passed=(
                fraction>=float(gate["minimum_nonmissing_fraction_of_geometry_qualified_focals"])
                and len(keep)>=int(gate["minimum_nonmissing_rows"])
                and ysd>float(gate["response_sd_gt"])
                and rank==X.shape[1]
                and cond<=float(gate["maximum_condition_number_after_missing_filter"])
            )
            if passed:
                yz=(ylog-ylog.mean())/ysd
                beta,_,brank,_=np.linalg.lstsq(X,yz,rcond=None)
                if int(brank)!=X.shape[1]:
                    raise RuntimeError(f"{plot}: rank failure during frozen response fit")
                coefs[plot]={
                    "target":float(beta[target_index]),
                    "direct_source_isolation":float(beta[direct_index]),
                    "topology_gain":float(beta[gain_index]),
                    "n":len(keep),
                }
        else:
            passed=False
        audit["pass"]=bool(passed)
        audits.append(audit)

    minplots=int(protocol["response_gate"]["study_pass"]["minimum_response_qualified_plots"])
    study_pass=len(coefs)>=minplots
    return coefs,{
        "plot_audits":audits,
        "qualified_plots":sorted(coefs,key=lambda x:int(x[2:])),
        "qualified_plot_count":len(coefs),
        "study_pass":bool(study_pass),
    }

def bootstrap_summary(coefs,protocol):
    plots=sorted(coefs,key=lambda x:int(x[2:]))
    vals=np.asarray([coefs[p]["target"] for p in plots],dtype=float)
    estimate=float(vals.mean())
    cfg=protocol["H1_primary"]["bootstrap"]
    rng=np.random.default_rng(int(cfg["seed"]))
    reps=int(cfg["replicates"])
    means=np.empty(reps,dtype=float)
    m=len(vals)
    for i in range(reps):
        idx=rng.integers(0,m,size=m)
        means[i]=vals[idx].mean()
    ci=np.quantile(means,[0.025,0.975],method="linear")
    loo={}
    for p in plots:
        other=np.asarray([coefs[q]["target"] for q in plots if q!=p],dtype=float)
        loo[p]=float(other.mean())
    direct=float(np.mean([coefs[p]["direct_source_isolation"] for p in plots]))
    gain=float(np.mean([coefs[p]["topology_gain"] for p in plots]))
    return {
        "estimate":estimate,
        "bootstrap_95_ci":[float(ci[0]),float(ci[1])],
        "prediction":"positive",
        "supported":bool(ci[0]>0),
        "qualified_plots":plots,
        "plot_coefficients":{p:coefs[p]["target"] for p in plots},
        "fraction_plot_coefficients_gt_zero":float(np.mean(vals>0)),
        "leave_one_plot_out_mean":loo,
        "bootstrap_estimates_sha256":sha([float(x) for x in means]),
    },{
        "mean_direct_source_isolation_coefficient":direct,
        "mean_topology_gain_coefficient":gain,
        "cannot_rescue_H1":True,
    }

def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--xlsx",type=Path,required=True)
    ap.add_argument("--protocol",type=Path,required=True)
    ap.add_argument("--authorization",type=Path,required=True)
    a=ap.parse_args()

    protocol=load_json(a.protocol)
    auth=load_json(a.authorization)
    if protocol.get("protocol_fingerprint_unsigned")!=EXPECTED_PROTOCOL:
        raise RuntimeError("protocol fingerprint drift")
    if auth.get("schema")!="structural.murundu_relative_topology_one_shot_authorization.v0_1":
        raise RuntimeError("unexpected authorization schema")
    if auth.get("status")!="AUTHORIZED_FOR_ONE_SHOT_WORKBOOK_ACCESS":
        raise RuntimeError("one-shot authorization absent")
    if auth.get("protocol_fingerprint")!=EXPECTED_PROTOCOL:
        raise RuntimeError("authorization/protocol drift")
    if auth.get("source",{}).get("doi")!="10.5061/dryad.612jm64kr":
        raise RuntimeError("Dryad DOI drift")
    if auth.get("source",{}).get("file_name")!="Ecological_Data.xlsx":
        raise RuntimeError("Dryad file-name drift")
    observed_sha=file_sha(a.xlsx)
    observed_bytes=a.xlsx.stat().st_size
    if observed_sha!=auth["source"]["sha256"]:
        raise RuntimeError("workbook SHA-256 drift")
    if observed_bytes!=int(auth["source"]["bytes"]):
        raise RuntimeError("workbook byte-size drift")

    wb=load_workbook(a.xlsx,read_only=True,data_only=True)
    sheet_names={s.lower():s for s in wb.sheetnames}
    sheet=sheet_names.get("ecological_data")
    if sheet is None:
        raise RuntimeError(f"Ecological_data worksheet absent: {wb.sheetnames}")
    ws=wb[sheet]
    header=[cell.value for cell in next(ws.iter_rows(min_row=1,max_row=1))]
    idx={str(name).strip():i+1 for i,name in enumerate(header) if name is not None}
    for col in (*GEOM_COLUMNS,PRIMARY_RESPONSE):
        if col not in idx:
            raise RuntimeError(f"required column absent: {col}")
    nrows=ws.max_row-1
    if nrows!=int(protocol["selection_provenance"]["declared_rows"]):
        raise RuntimeError(f"row count drift: {nrows}")

    # Geometry-only stage. Richness.Trees is deliberately not read here.
    columns={c:col_values(ws,idx[c],nrows) for c in GEOM_COLUMNS}
    records=[]
    for i in range(nrows):
        records.append({
            "row_index":i,
            "id":"" if columns["ID.murundu"][i] is None else str(columns["ID.murundu"][i]),
            "plot":"" if columns["Plots"][i] is None else str(columns["Plots"][i]).strip(),
            "lon":number(columns["X"][i]),
            "lat":number(columns["Y"][i]),
            "area":number(columns["Area_Mur_m2"][i]),
            "height":number(columns["Height_Mur"][i]),
        })

    qualified,geometry=geometry_stage(records,protocol)
    base={
        "schema":"structural.murundu_relative_topology_outcome.v0_1",
        "source":{
            "doi":"10.5061/dryad.612jm64kr",
            "file_name":"Ecological_Data.xlsx",
            "sha256":observed_sha,
            "bytes":observed_bytes,
        },
        "protocol_fingerprint":EXPECTED_PROTOCOL,
        "response_access":{
            "workbook_opened":True,
            "geometry_columns_parsed":list(GEOM_COLUMNS),
            "primary_response_column_parsed":False,
            "herb_response_parsed":False,
            "termite_response_parsed":False,
        },
        "geometry_gate":geometry,
        "claim_boundary":protocol["claim_boundary"],
    }
    if not geometry["study_pass"]:
        base.update({
            "status":"TERMINAL_STOP_GEOMETRY_NON_ESTIMABLE",
            "response_gate":None,
            "H1_primary":None,
            "secondary_nonrescuing":None,
        })
        base["outcome_fingerprint"]=sha(base)
        print(json.dumps(base,indent=2,sort_keys=True))
        return 2

    # Only after the response-blind geometry gate passes do we parse the one
    # frozen primary response column.
    response_values=col_values(ws,idx[PRIMARY_RESPONSE],nrows)
    base["response_access"]["primary_response_column_parsed"]=True
    coefs,response=response_stage(qualified,response_values,protocol)
    base["response_gate"]=response
    if not response["study_pass"]:
        base.update({
            "status":"TERMINAL_STOP_RESPONSE_DESIGN_NON_ESTIMABLE",
            "H1_primary":None,
            "secondary_nonrescuing":None,
        })
        base["outcome_fingerprint"]=sha(base)
        print(json.dumps(base,indent=2,sort_keys=True))
        return 2

    h1,secondary=bootstrap_summary(coefs,protocol)
    base.update({
        "status":"ONE_SHOT_TOPOLOGY_EXTENSION_SCORED",
        "H1_primary":h1,
        "secondary_nonrescuing":secondary,
        "evidence_role":"prospective topology extension into habitat fragmentation; does not count as independent Structural confirmation",
    })
    base["outcome_fingerprint"]=sha(base)
    print(json.dumps(base,indent=2,sort_keys=True))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
