#!/usr/bin/env python3
"""Freeze the Roeble et al. 2024 island-origin crosswalk before pilot response."""
from __future__ import annotations

from collections import Counter, defaultdict
import argparse
import hashlib
import io
import json
import math
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

URL = "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-024-51556-7/MediaObjects/41467_2024_51556_MOESM6_ESM.xlsx"
EXPECTED_SHA256 = "fdca6086440c721e2766c3b73b96969c401b1a267a0cf9590c912d369d5e8597"
DOI = "10.1038/s41467-024-51556-7"
MIN_ENTITY_MATCHES = 5
MIN_ENTITY_COVERAGE = 0.50
MIN_DOMINANT_ARCHIPELAGO_SHARE = 0.80
MIN_QUANTITATIVE_TYPE_COVERAGE = 0.80

def sha(v) -> str:
    return hashlib.sha256(json.dumps(v,sort_keys=True,separators=(",",":"),ensure_ascii=False).encode()).hexdigest()

def fetch_bytes(url):
    req=Request(url,headers={"User-Agent":"Structural-GIFT-Roeble-crosswalk-freeze/0.1"})
    with urlopen(req,timeout=180) as r:
        return r.read()

def norm_type(x):
    if x is None:
        return None
    t=str(x).strip().lower()
    if t in {"continental","oceanic","mixed"}:
        return t
    return None

def main() -> int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--universe",type=Path,required=True)
    args=ap.parse_args()

    u=json.loads(args.universe.read_text(encoding="utf-8"))
    if u.get("response_values_accessed") is not False:
        raise RuntimeError("universe response already accessed")

    data=fetch_bytes(URL)
    digest=hashlib.sha256(data).hexdigest()
    if digest != EXPECTED_SHA256:
        raise RuntimeError(f"Roeble Supplementary Data 3 SHA drift: {digest}")

    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    ws=wb[wb.sheetnames[0]]
    rows=ws.iter_rows(values_only=True)
    header=[str(x) if x is not None else "" for x in next(rows)]
    required=["entity_ID","geo_entity","archipelago","geology","island_type"]
    missing=[x for x in required if x not in header]
    if missing:
        raise RuntimeError("missing Roeble columns: "+",".join(missing))
    ix={name:header.index(name) for name in required}

    raw=[]
    for row in rows:
        entity=row[ix["entity_ID"]]
        if entity is None:
            continue
        try:
            entity_id=str(int(float(entity)))
        except (TypeError,ValueError):
            continue
        raw.append({
            "entity_ID":entity_id,
            "geo_entity":None if row[ix["geo_entity"]] is None else str(row[ix["geo_entity"]]).strip(),
            "archipelago":None if row[ix["archipelago"]] is None else str(row[ix["archipelago"]]).strip(),
            "geology":None if row[ix["geology"]] is None else str(row[ix["geology"]]).strip(),
            "island_type":norm_type(row[ix["island_type"]]),
        })

    by_entity=defaultdict(list)
    by_roeble_arch=defaultdict(list)
    for row in raw:
        by_entity[row["entity_ID"]].append(row)
        if row["archipelago"]:
            by_roeble_arch[row["archipelago"]].append(row)

    duplicate_conflicts=[]
    canonical={}
    for eid, vals in by_entity.items():
        pairs={(v["archipelago"],v["island_type"]) for v in vals}
        if len(pairs)>1:
            duplicate_conflicts.append({"entity_ID":eid,"pairs":sorted([list(p) for p in pairs],key=str)})
            continue
        canonical[eid]=vals[0]

    roeble_arch_type={}
    for name, vals in by_roeble_arch.items():
        types=[v["island_type"] for v in vals if v["island_type"] in {"continental","oceanic"}]
        c=Counter(types)
        roeble_arch_type[name]={
            "n_rows":len(vals),
            "typed_rows":sum(c.values()),
            "continental":c["continental"],
            "oceanic":c["oceanic"],
            "oceanic_fraction":(c["oceanic"]/sum(c.values()) if c else None),
            "class":(
                "continental" if c and c["oceanic"]==0
                else "oceanic" if c and c["continental"]==0
                else "mixed" if c else "unknown"
            ),
        }

    confirm=set(u["confirmatory_archipelagos"])
    out=[]
    for g in u["groups"]:
        ids=[str(i["entity_ID"]) for i in g["islands"]]
        matched=[canonical[eid] for eid in ids if eid in canonical]
        names=Counter(r["archipelago"] for r in matched if r["archipelago"])
        dominant=None
        dominant_n=0
        if names:
            dominant,dominant_n=sorted(names.items(),key=lambda kv:(-kv[1],kv[0]))[0]
        matched_fraction=len(matched)/len(ids)
        dominant_share=(dominant_n/len(matched) if matched else 0.0)
        mapped=(
            dominant is not None
            and len(matched)>=MIN_ENTITY_MATCHES
            and matched_fraction>=MIN_ENTITY_COVERAGE
            and dominant_share>=MIN_DOMINANT_ARCHIPELAGO_SHARE
        )
        external=roeble_arch_type.get(dominant) if mapped else None

        island_types=[r["island_type"] for r in matched if r["island_type"] in {"continental","oceanic"}]
        type_coverage=len(island_types)/len(ids)
        oceanic_fraction=(
            sum(t=="oceanic" for t in island_types)/len(island_types)
            if island_types and type_coverage>=MIN_QUANTITATIVE_TYPE_COVERAGE
            else None
        )
        out.append({
            "archipelago_id":g["archipelago_id"],
            "partition":"confirmatory" if g["archipelago_id"] in confirm else "pilot",
            "n_islands":len(ids),
            "entity_matches":len(matched),
            "entity_match_fraction":matched_fraction,
            "dominant_roeble_archipelago":dominant,
            "dominant_roeble_archipelago_n":dominant_n,
            "dominant_roeble_archipelago_share":dominant_share,
            "roeble_archipelago_mapping_qualified":mapped,
            "roeble_archipelago_type":external["class"] if external else None,
            "roeble_archipelago_oceanic_fraction":external["oceanic_fraction"] if external else None,
            "island_type_rows":len(island_types),
            "island_type_coverage":type_coverage,
            "direct_oceanic_fraction":oceanic_fraction,
        })

    conf=[r for r in out if r["partition"]=="confirmatory"]
    q=[r for r in conf if r["direct_oceanic_fraction"] is not None]
    pure_c=sum(r["direct_oceanic_fraction"]==0 for r in q)
    pure_o=sum(r["direct_oceanic_fraction"]==1 for r in q)
    mixed=sum(0<r["direct_oceanic_fraction"]<1 for r in q)

    payload={
        "schema":"structural.gift_roeble_geology_crosswalk.v0_1",
        "status":"FROZEN_BEFORE_PILOT_RESPONSE",
        "source":{
            "doi":DOI,
            "supplement":"Supplementary Data 3",
            "url":URL,
            "xlsx_sha256":digest,
            "xlsx_bytes":len(data),
        },
        "response_values_accessed":False,
        "gift_species_composition_accessed":False,
        "matching_rule":{
            "key":"GIFT entity_ID",
            "minimum_entity_matches":MIN_ENTITY_MATCHES,
            "minimum_entity_coverage":MIN_ENTITY_COVERAGE,
            "minimum_dominant_roeble_archipelago_share":MIN_DOMINANT_ARCHIPELAGO_SHARE,
            "minimum_quantitative_island_type_coverage":MIN_QUANTITATIVE_TYPE_COVERAGE,
            "no_name_fuzzy_matching":True,
            "no_response_based_resolution":True,
        },
        "duplicate_entity_conflicts":duplicate_conflicts,
        "archipelagos":out,
        "confirmatory_quantitative_geology_eligible":len(q),
        "confirmatory_pure_continental":pure_c,
        "confirmatory_pure_oceanic":pure_o,
        "confirmatory_mixed":mixed,
        "geology_h2_class_contrast_estimable":pure_c>=3 and pure_o>=3,
        "geology_h2_fraction_slope_estimable":len(q)>=6 and min(r["direct_oceanic_fraction"] for r in q)<max(r["direct_oceanic_fraction"] for r in q),
        "claim_boundary":"external island-type classification is a geological-history moderator; it does not identify colonization, rescue, persistence or demographic mechanism",
    }
    payload["crosswalk_fingerprint"]=sha({
        "source_sha256":digest,
        "matching_rule":payload["matching_rule"],
        "archipelagos":out,
    })
    print(json.dumps(payload,indent=2,sort_keys=True))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
