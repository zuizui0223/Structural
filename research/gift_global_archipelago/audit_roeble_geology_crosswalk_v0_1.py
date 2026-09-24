#!/usr/bin/env python3
"""Audit Roeble et al. 2024 Supplementary Data 3 as a response-blind geology crosswalk source."""
from __future__ import annotations

import hashlib
import io
import json
from urllib.request import Request, urlopen

from openpyxl import load_workbook

URL = "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-024-51556-7/MediaObjects/41467_2024_51556_MOESM6_ESM.xlsx"
DOI = "10.1038/s41467-024-51556-7"

def fetch_bytes(url: str) -> bytes:
    req=Request(url,headers={"User-Agent":"Structural-GIFT-geology-crosswalk-audit/0.1"})
    with urlopen(req,timeout=180) as r:
        return r.read()

def main() -> int:
    data=fetch_bytes(URL)
    digest=hashlib.sha256(data).hexdigest()
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    sheets=[]
    for ws in wb.worksheets:
        preview=[]
        for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,12),values_only=True):
            preview.append([None if x is None else str(x) for x in row[:40]])
        sheets.append({
            "title":ws.title,
            "rows":ws.max_row,
            "columns":ws.max_column,
            "preview":preview,
        })
    payload={
        "schema":"structural.roeble_2024_geology_crosswalk_audit.v0_1",
        "status":"response_blind_external_metadata_audit",
        "doi":DOI,
        "supplement":"Supplementary Data 3",
        "url":URL,
        "xlsx_bytes":len(data),
        "xlsx_sha256":digest,
        "gift_species_composition_accessed":False,
        "structural_outcome_accessed":False,
        "sheets":sheets,
    }
    print(json.dumps(payload,indent=2,sort_keys=True))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
