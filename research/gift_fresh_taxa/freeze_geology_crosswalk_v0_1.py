#!/usr/bin/env python3
"""Freeze Roeble et al. geology crosswalk for a fresh taxonomic GIFT protocol."""
from __future__ import annotations
from collections import Counter,defaultdict
import argparse,hashlib,io,json
from pathlib import Path
from urllib.request import Request,urlopen
from openpyxl import load_workbook

URL="https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-024-51556-7/MediaObjects/41467_2024_51556_MOESM6_ESM.xlsx"
EXPECTED="fdca6086440c721e2766c3b73b96969c401b1a267a0cf9590c912d369d5e8597"
DOI="10.1038/s41467-024-51556-7"
def sha(x):return hashlib.sha256(json.dumps(x,sort_keys=True,separators=(",",":"),ensure_ascii=False).encode()).hexdigest()
def fetch():
    req=Request(URL,headers={"User-Agent":"Structural-GIFT-fresh-geology/0.1"})
    with urlopen(req,timeout=180) as r:return r.read()
def norm(x):
    if x is None:return None
    t=str(x).strip().lower()
    return t if t in {"continental","oceanic","mixed"} else None
def main():
    ap=argparse.ArgumentParser();ap.add_argument("--universe",type=Path,required=True);args=ap.parse_args()
    u=json.loads(args.universe.read_text())
    if u["response_values_accessed"] is not False:raise RuntimeError("response already opened")
    data=fetch();digest=hashlib.sha256(data).hexdigest()
    if digest!=EXPECTED:raise RuntimeError("Roeble supplement drift")
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True);ws=wb[wb.sheetnames[0]]
    it=ws.iter_rows(values_only=True);header=[str(x) if x is not None else "" for x in next(it)]
    need=["entity_ID","archipelago","island_type"];ix={k:header.index(k) for k in need}
    by=defaultdict(list)
    for row in it:
        try:eid=str(int(float(row[ix["entity_ID"]])))
        except (TypeError,ValueError):continue
        by[eid].append({"archipelago":None if row[ix["archipelago"]] is None else str(row[ix["archipelago"]]).strip(),"island_type":norm(row[ix["island_type"]])})
    canonical={}
    for eid,vals in by.items():
        pairs={(v["archipelago"],v["island_type"]) for v in vals}
        if len(pairs)==1:canonical[eid]=vals[0]
    rows=[]
    for g in u["groups"]:
        ids=[]
        # Universe stores list IDs but not island rows; membership fingerprint only.
        # Reconstruct direct entity mapping is impossible here, so require a later
        # universe schema that retains entity IDs. Fail closed instead of fuzzy matching.
        if "entity_ids" not in g:
            raise RuntimeError("fresh universe must retain entity_ids before geology crosswalk can be frozen")
        ids=[str(x) for x in g["entity_ids"]]
        vals=[canonical[eid] for eid in ids if eid in canonical and canonical[eid]["island_type"] in {"continental","oceanic"}]
        frac=(sum(v["island_type"]=="oceanic" for v in vals)/len(vals)) if vals and len(vals)/len(ids)>=0.80 else None
        rows.append({"archipelago_id":g["archipelago_id"],"n_islands":len(ids),"typed_islands":len(vals),"direct_oceanic_fraction":frac})
    q=[r for r in rows if r["direct_oceanic_fraction"] is not None]
    out={"schema":"structural.gift_fresh_taxon_geology.v0_1","status":"FROZEN_BEFORE_PILOT_RESPONSE","taxon_name":u["taxon_name"],"universe_fingerprint":u["universe_fingerprint"],"source":{"doi":DOI,"xlsx_sha256":digest,"url":URL},"response_values_accessed":False,"gift_species_composition_accessed":False,"archipelagos":rows,"quantitative_geology_eligible_all":len(q),"claim_boundary":"geological-history moderator only; not a demographic mechanism"}
    out["crosswalk_fingerprint"]=sha(out);print(json.dumps(out,indent=2,sort_keys=True));return 0
if __name__=="__main__":raise SystemExit(main())
